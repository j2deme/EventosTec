"""Tests para GET /api/activities/batch-template (plantilla XLSX descargable).

Cubre:
- Descarga sin autenticación (el <a download> del modal no envía Authorization).
- Estructura del libro: hojas, encabezados y filas de ejemplo.
- Instrucciones documentan todas las columnas de BATCH_COLUMNS.
- Las filas de ejemplo pasan el dry-run del importador (fuente única compartida).
"""

from io import BytesIO

from openpyxl import load_workbook

from app.services.activity_service import (
    BATCH_COLUMNS,
    build_activities_xlsx_template,
    create_activities_from_xlsx,
)


def _headers():
    return [c["header"] for c in BATCH_COLUMNS]


def test_batch_template_returns_attachment_without_auth(client):
    """El endpoint no exige JWT: solo expone estructura, sin datos reales."""
    resp = client.get("/api/activities/batch-template")

    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.content_type
    disposition = resp.headers.get("Content-Disposition", "")
    assert "attachment" in disposition
    assert "plantilla_actividades.xlsx" in disposition


def test_batch_template_headers_match_batch_columns(client):
    """El libro tiene las hojas esperadas y los encabezados de BATCH_COLUMNS."""
    resp = client.get("/api/activities/batch-template")
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Plantilla", "Instrucciones"]

    ws = wb["Plantilla"]
    headers = [cell.value for cell in ws[1]]
    assert headers == _headers()
    # Encabezado + 2 filas de ejemplo
    assert ws.max_row >= 3


def test_batch_template_documents_all_columns(client):
    """La hoja Instrucciones lista cada columna con su obligatoriedad."""
    resp = client.get("/api/activities/batch-template")
    wb = load_workbook(BytesIO(resp.data))

    wi = wb["Instrucciones"]
    rows = [list(r) for r in wi.iter_rows(min_row=2, values_only=True)]
    by_column = {r[0]: r for r in rows if r and r[0]}

    for col in BATCH_COLUMNS:
        assert col["header"] in by_column, f"columna sin documentar: {col['header']}"
        expected_flag = "Sí" if col["required"] else "No"
        assert by_column[col["header"]][1] == expected_flag


def test_template_example_rows_pass_dry_run_import():
    """Las filas de ejemplo son importables: dry-run sin inválidas.

    Garantiza que la plantilla nunca envíe ejemplos que el esquema rechace
    (mismos validadores que POST /api/activities/batch).
    """
    buf = build_activities_xlsx_template()

    report = create_activities_from_xlsx(buf, event_id=1, dry_run=True)

    assert report["summary"]["invalid"] == 0, report["errors"]
    assert report["summary"]["valid"] == 2
    assert report["created"] == 0
