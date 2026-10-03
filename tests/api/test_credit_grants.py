"""Tests de la Fase 3: tabla ``credit_grants`` (historial de otorgamientos).

Cubre POST/GET /api/students/credit-grants y su efecto en
GET /api/students/complementary-credits y en el Excel:

- acreditar "gasta" los eventos y el estudiante deja de aparecer;
- con horas frescas en otros eventos acredita OTRO crédito;
- el Excel nunca incluye a los ya otorgados (es lo que se sube al externo);
- el historial agrupa los grants por lote (batch_id).
"""

from datetime import datetime, timedelta

from app import db
from app.models.activity import Activity
from app.models.event import Event
from app.models.registration import Registration
from app.models.student import Student


# ---------------------------------------------------------------------------
# helpers locales
# ---------------------------------------------------------------------------


def _create_event(app, name, days_offset=0):
    """Crea un evento desplazado en el tiempo (para fijar el orden cronológico)."""
    with app.app_context():
        start = datetime.now() + timedelta(days=days_offset)
        event = Event(
            name=name,
            description="Test credit_grants",
            start_date=start,
            end_date=start + timedelta(days=7),
            is_active=True,
        )
        db.session.add(event)
        db.session.commit()
        return event.id


def _create_student(app, control, name, career="Ingeniería en Sistemas"):
    with app.app_context():
        student = Student(
            control_number=control,
            full_name=name,
            career=career,
            email=f"{control.lower()}@test.com",
        )
        db.session.add(student)
        db.session.commit()
        return student.id


def _add_participation(app, student_id, event_id, hours, name="Act Grants"):
    with app.app_context():
        activity = Activity(
            event_id=event_id,
            department="TEST",
            name=name,
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=hours),
            duration_hours=hours,
            activity_type="Conferencia",
            location="Test",
            modality="Presencial",
        )
        db.session.add(activity)
        db.session.flush()
        db.session.add(
            Registration(
                student_id=student_id, activity_id=activity.id, status="Asistió"
            )
        )
        db.session.commit()
        return activity.id


def _list_credits(client, auth_headers, event_ids, extra=""):
    ids = ",".join(str(e) for e in event_ids)
    return client.get(
        f"/api/students/complementary-credits?event_ids={ids}{extra}",
        headers=auth_headers,
    )


def _grant(client, auth_headers, event_ids, **extra):
    payload = {"event_ids": list(event_ids), "confirm": True}
    payload.update(extra)
    return client.post(
        "/api/students/credit-grants", json=payload, headers=auth_headers
    )


def _grant_rows(app):
    """(student_id, event_id, batch_id, granted_by) de toda la tabla."""
    with app.app_context():
        from app.models.credit_grant import CreditGrant

        return [
            (g.student_id, g.event_id, g.batch_id, g.granted_by)
            for g in CreditGrant.query.order_by(CreditGrant.id).all()
        ]


# ---------------------------------------------------------------------------
# contrato del endpoint de otorgamiento
# ---------------------------------------------------------------------------


def test_credit_grants_require_auth(client):
    assert client.get("/api/students/credit-grants").status_code == 401
    assert (
        client.post("/api/students/credit-grants", json={"confirm": True}).status_code
        == 401
    )


def test_grant_validations(app, client, auth_headers):
    ev = _create_event(app, "Grants Validación")
    sid = _create_student(app, "GRV01", "Alumno Validación")
    _add_participation(app, sid, ev, 12.0)

    # Sin confirm: la operación no se ejecuta (confirmación explícita)
    resp = client.post(
        "/api/students/credit-grants",
        json={"event_ids": [ev]},
        headers=auth_headers,
    )
    assert resp.status_code == 400
    assert "confirm" in resp.get_json()["message"]

    # Sin event_ids
    resp = client.post(
        "/api/students/credit-grants", json={"confirm": True}, headers=auth_headers
    )
    assert resp.status_code == 400
    assert "event_ids" in resp.get_json()["message"]

    # event_ids vacío / no numérico
    resp = client.post(
        "/api/students/credit-grants",
        json={"confirm": True, "event_ids": []},
        headers=auth_headers,
    )
    assert resp.status_code == 400

    resp = client.post(
        "/api/students/credit-grants",
        json={"confirm": True, "event_ids": ["abc"]},
        headers=auth_headers,
    )
    assert resp.status_code == 400

    # Evento inexistente
    resp = client.post(
        "/api/students/credit-grants",
        json={"confirm": True, "event_ids": [99999]},
        headers=auth_headers,
    )
    assert resp.status_code == 404

    # note demasiado largo
    resp = _grant(client, auth_headers, [ev], note="x" * 256)
    assert resp.status_code == 400

    # Nada quedó guardado y el estudiante sigue pendiente
    assert _grant_rows(app) == []
    resp = _list_credits(client, auth_headers, [ev])
    assert resp.get_json()["total_students"] == 1


# ---------------------------------------------------------------------------
# efecto en la lista de créditos
# ---------------------------------------------------------------------------


def test_grant_records_and_then_excludes_from_list(app, client, auth_headers):
    ev = _create_event(app, "Grants Único")
    sid = _create_student(app, "GR010", "Alumno Otorgado")
    _add_participation(app, sid, ev, 12.0)

    # Antes de otorgar: aparece y nada está gastado
    resp = _list_credits(client, auth_headers, [ev])
    data = resp.get_json()
    assert data["total_students"] == 1
    assert data["excluded_already_granted"] == 0
    student = data["students"][0]
    assert student["hours_consumed"] == 0.0
    assert student["hours_available"] == 12.0
    assert student["crossing_event_id"] == ev
    assert student["granted_event_ids"] == []
    assert student["already_granted"] is False

    # Otorgar
    resp = _grant(client, auth_headers, [ev])
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["granted_students"] == 1
    assert body["granted_events"] == 1
    assert body["batch_id"]

    # Una fila en credit_grants, con el evento gastado y el admin que confirmó
    rows = _grant_rows(app)
    assert len(rows) == 1
    assert rows[0][0] == sid
    assert rows[0][1] == ev
    assert rows[0][2] == body["batch_id"]
    assert rows[0][3] is not None

    # Ya no aparece en la lista
    resp = _list_credits(client, auth_headers, [ev])
    data = resp.get_json()
    assert data["total_students"] == 0
    assert data["excluded_already_granted"] == 1
    assert data["excluded_already_credited"] == 0

    # La vista de auditoría lo recupera con la marca
    resp = _list_credits(client, auth_headers, [ev], extra="&include_granted=1")
    data = resp.get_json()
    assert data["total_students"] == 1
    student = data["students"][0]
    assert student["already_granted"] is True
    assert student["granted_event_ids"] == [ev]
    assert student["hours_consumed"] == 12.0
    assert student["hours_available"] == 0.0


def test_second_credit_requires_fresh_hours(app, client, auth_headers):
    """Tras acreditar en A, otras 10 h en B liberan OTRO crédito."""
    ev1 = _create_event(app, "Grants Segundo A", days_offset=0)
    ev2 = _create_event(app, "Grants Segundo B", days_offset=3)
    sid = _create_student(app, "GR020", "Alumno Segundo Crédito")
    _add_participation(app, sid, ev1, 12.0)
    _add_participation(app, sid, ev2, 10.0)

    # Acreditado en A: A+B muestra solo las horas frescas de B
    resp = _grant(client, auth_headers, [ev1])
    assert resp.status_code == 200

    resp = _list_credits(client, auth_headers, [ev1, ev2])
    data = resp.get_json()
    assert data["total_students"] == 1
    student = data["students"][0]
    assert student["hours_consumed"] == 12.0
    assert student["hours_available"] == 10.0
    assert student["crossing_event_id"] == ev2
    assert student["granted_event_ids"] == [ev1]

    # El segundo lote gasta SOLO B (A ya estaba gastado)
    resp = _grant(client, auth_headers, [ev1, ev2])
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["granted_students"] == 1
    assert body["granted_events"] == 1

    rows = _grant_rows(app)
    assert sorted((r[0], r[1]) for r in rows) == [(sid, ev1), (sid, ev2)]
    assert rows[0][2] != rows[1][2]  # dos lotes distintos

    # Sin horas frescas: queda excluido como ya otorgado
    resp = _list_credits(client, auth_headers, [ev1, ev2])
    data = resp.get_json()
    assert data["total_students"] == 0
    assert data["excluded_already_granted"] == 1


def test_grant_is_not_repeatable(app, client, auth_headers):
    """Repetir la acción no vuelve a gastar ni a insertar filas."""
    ev = _create_event(app, "Grants Idempotente")
    sid = _create_student(app, "GR030", "Alumno Repetido")
    _add_participation(app, sid, ev, 12.0)

    resp = _grant(client, auth_headers, [ev])
    assert resp.status_code == 200

    resp = _grant(client, auth_headers, [ev])
    assert resp.status_code == 400
    assert "pendientes" in resp.get_json()["message"]

    assert len(_grant_rows(app)) == 1


def test_override_forced_row_only_consumes_events_with_hours(app, client, auth_headers):
    """Una fila forzada por override 'include' gasta solo lo que aportó horas."""
    ev1 = _create_event(app, "Grants Override A", days_offset=0)
    ev2 = _create_event(app, "Grants Override B", days_offset=3)
    sid = _create_student(app, "GR040", "Alumno Forzado")
    _add_participation(app, sid, ev1, 12.0)
    _add_participation(app, sid, ev2, 5.0)

    # Cruza en A (evento anterior al último) → lo excluye la regla derivada...
    resp = _list_credits(client, auth_headers, [ev1, ev2])
    assert resp.get_json()["total_students"] == 0

    # ... salvo el override 'include'
    resp = client.post(
        "/api/students/credit-overrides",
        json={"student_id": sid, "decision": "include"},
        headers=auth_headers,
    )
    assert resp.status_code == 200

    resp = _list_credits(client, auth_headers, [ev1, ev2])
    data = resp.get_json()
    assert data["total_students"] == 1
    # El cruce es en A: solo A se gastaría (B queda para otro crédito)
    assert data["students"][0]["crossing_event_id"] == ev1
    assert data["students"][0]["grant_event_ids"] == [ev1]

    resp = _grant(client, auth_headers, [ev1, ev2])
    assert resp.status_code == 200
    assert resp.get_json()["granted_events"] == 1

    rows = _grant_rows(app)
    assert [(r[0], r[1]) for r in rows] == [(sid, ev1)]

    # El override sigue forzando su presencia en la lista
    resp = _list_credits(client, auth_headers, [ev1, ev2])
    data = resp.get_json()
    assert data["total_students"] == 1
    assert data["students"][0]["granted_event_ids"] == [ev1]


# ---------------------------------------------------------------------------
# Excel e historial
# ---------------------------------------------------------------------------


def test_export_never_includes_already_granted(app, client, auth_headers):
    """El XLSX excluye a los acreditados y deja constancia en el resumen."""
    from io import BytesIO

    from openpyxl import load_workbook

    ev = _create_event(app, "Grants Export")
    sid = _create_student(app, "GR050", "Alumno Export")
    _add_participation(app, sid, ev, 12.0)

    assert _grant(client, auth_headers, [ev]).status_code == 200

    resp = client.get(
        f"/api/students/complementary-credits/export?event_ids={ev}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert (
        resp.content_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    # Sin filas de datos: el estudiante ya fue acreditado
    assert ws.cell(row=5, column=2).value is None
    # El resumen documenta por qué no está
    summary = ws.cell(row=6, column=1).value
    assert summary.startswith("Total de estudiantes: 0")
    assert "ya otorgados" in summary


def test_get_credit_grants_history(app, client, auth_headers):
    ev1 = _create_event(app, "Grants Histórico A", days_offset=0)
    ev2 = _create_event(app, "Grants Histórico B", days_offset=3)
    sid = _create_student(app, "GR060", "Alumno Historial")
    _add_participation(app, sid, ev1, 12.0)
    _add_participation(app, sid, ev2, 10.0)

    assert _grant(client, auth_headers, [ev1]).status_code == 200
    assert _grant(client, auth_headers, [ev2]).status_code == 200

    resp = client.get("/api/students/credit-grants", headers=auth_headers)
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["total"] == 2

    batches = {b["batch_id"]: b for b in data["batches"]}
    assert len(batches) == 2

    first = next(b for b in data["batches"] if b["event_ids"] == [ev1])
    assert first["student_ids"] == [sid]
    assert first["student_count"] == 1
    assert first["event_names"] == ["Grants Histórico A"]
    assert first["granted_at"]
    assert first["granted_by"] is not None
