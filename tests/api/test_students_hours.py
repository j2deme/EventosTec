"""Tests for student hours tracking endpoints."""

import pytest
from datetime import datetime, timedelta


@pytest.fixture
def sample_event(app):
    """Create a sample event for testing."""
    from app.models.event import Event
    from app import db

    with app.app_context():
        event = Event(
            name="Test Event",
            description="Test event for hours tracking",
            start_date=datetime.now(),
            end_date=datetime.now() + timedelta(days=7),
            is_active=True,
        )
        db.session.add(event)
        db.session.commit()
        event_id = event.id

    return event_id


@pytest.fixture
def sample_student(app):
    """Create a sample student for testing."""
    from app.models.student import Student
    from app import db

    with app.app_context():
        student = Student(
            control_number="TEST001",
            full_name="Test Student",
            career="Test Career",
            email="test@test.com",
        )
        db.session.add(student)
        db.session.commit()
        student_id = student.id

    return student_id


@pytest.fixture
def sample_activity(app, sample_event):
    """Create a sample activity for testing."""
    from app.models.activity import Activity
    from app import db

    with app.app_context():
        activity = Activity(
            event_id=sample_event,
            department="TEST",
            name="Test Activity",
            description="Test activity",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=2),
            duration_hours=2.0,
            activity_type="Conferencia",
            location="Test Location",
            modality="Presencial",
        )
        db.session.add(activity)
        db.session.commit()
        activity_id = activity.id

    return activity_id


@pytest.fixture
def sample_registration(app, sample_student, sample_activity):
    """Create a sample registration with 'Asistió' status."""
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        registration = Registration(
            student_id=sample_student, activity_id=sample_activity, status="Asistió"
        )
        db.session.add(registration)
        db.session.commit()
        registration_id = registration.id

    return registration_id


def test_get_students_with_filters(
    client, sample_event, sample_student, sample_activity, sample_registration
):
    """Test getting students with event filter."""
    # This endpoint doesn't require auth for basic GET
    resp = client.get(f"/api/students/?event_id={sample_event}")
    assert resp.status_code == 200
    data = resp.get_json()
    assert "students" in data
    assert len(data["students"]) > 0


def test_get_student_hours_by_event(
    client, sample_student, sample_event, sample_registration
):
    """Test getting student hours grouped by event."""
    resp = client.get(f"/api/students/{sample_student}/hours-by-event")
    assert resp.status_code == 200
    data = resp.get_json()
    assert "events_hours" in data
    assert "student" in data

    # Should have at least one event with hours
    if len(data["events_hours"]) > 0:
        event_hours = data["events_hours"][0]
        assert "event_id" in event_hours
        assert "total_hours" in event_hours
        assert "has_complementary_credit" in event_hours
        assert event_hours["total_hours"] >= 0


def test_get_student_event_details(
    client, sample_student, sample_event, sample_registration
):
    """Test getting detailed student participation in an event."""
    resp = client.get(f"/api/students/{sample_student}/event/{sample_event}/details")
    assert resp.status_code == 200
    data = resp.get_json()
    assert "student" in data
    assert "event" in data
    assert "total_confirmed_hours" in data
    assert "has_complementary_credit" in data
    assert "activities" in data

    # Should have confirmation for hours >= 10
    if data["total_confirmed_hours"] >= 10.0:
        assert data["has_complementary_credit"] is True
    else:
        assert data["has_complementary_credit"] is False


def test_student_hours_counts_confirmed_and_attended(
    app, client, sample_student, sample_event
):
    """Test de la semántica unificada: 'Confirmado' y 'Asistió' cuentan;
    'Registrado' no (fuente única: app/services/hours_service.py)."""
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        # Create activities with different statuses
        activities_data = [
            ("Asistió", 5.0),
            ("Confirmado", 3.0),
            ("Registrado", 2.0),
            ("Asistió", 6.0),
        ]

        for status, hours in activities_data:
            activity = Activity(
                event_id=sample_event,
                department="TEST",
                name=f"Activity {status}",
                description="Test",
                start_datetime=datetime.now(),
                end_datetime=datetime.now() + timedelta(hours=hours),
                duration_hours=hours,
                activity_type="Conferencia",
                location="Test",
                modality="Presencial",
            )
            db.session.add(activity)
            db.session.flush()

            registration = Registration(
                student_id=sample_student, activity_id=activity.id, status=status
            )
            db.session.add(registration)

        db.session.commit()

    # Get hours by event
    resp = client.get(f"/api/students/{sample_student}/hours-by-event")
    assert resp.status_code == 200
    data = resp.get_json()

    # Find the test event
    test_event = next(
        (e for e in data["events_hours"] if e["event_id"] == sample_event), None
    )
    assert test_event is not None

    # Semántica unificada: 'Asistió' (5.0 + 6.0) + 'Confirmado' (3.0) = 14.0;
    # 'Registrado' (2.0) no cuenta.
    assert test_event["total_hours"] == 14.0
    assert test_event["has_complementary_credit"] is True


def test_complementary_credit_badge_logic(app, client, sample_student, sample_event):
    """Test that complementary credit badge appears correctly."""
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        # Create activity with exactly 10 hours
        activity = Activity(
            event_id=sample_event,
            department="TEST",
            name="10 Hour Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=10),
            duration_hours=10.0,
            activity_type="Curso",
            location="Test",
            modality="Presencial",
        )
        db.session.add(activity)
        db.session.flush()

        registration = Registration(
            student_id=sample_student, activity_id=activity.id, status="Asistió"
        )
        db.session.add(registration)
        db.session.commit()

    resp = client.get(f"/api/students/{sample_student}/hours-by-event")
    assert resp.status_code == 200
    data = resp.get_json()

    test_event = next(
        (e for e in data["events_hours"] if e["event_id"] == sample_event), None
    )
    assert test_event is not None
    assert test_event["total_hours"] >= 10.0
    assert test_event["has_complementary_credit"] is True


def test_get_student_hours_not_found(client):
    """Test getting hours for non-existent student returns 404."""
    resp = client.get("/api/students/99999/hours-by-event")
    assert resp.status_code == 404
    data = resp.get_json()
    assert "message" in data


def test_get_student_event_details_not_found(client, sample_student):
    """Test getting event details for non-existent event returns 404."""
    resp = client.get(f"/api/students/{sample_student}/event/99999/details")
    assert resp.status_code == 404
    data = resp.get_json()
    assert "message" in data


def test_get_complementary_credits_requires_event_id(client, auth_headers):
    """Test that complementary credits endpoint requires event_id."""
    resp = client.get("/api/students/complementary-credits", headers=auth_headers)
    assert resp.status_code == 400
    data = resp.get_json()
    assert "event_id" in data["message"].lower()


def test_get_complementary_credits_filters_by_hours(
    app, client, auth_headers, sample_event
):
    """Test that only students with 10+ hours are returned."""
    from app.models.student import Student
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db
    from datetime import datetime, timedelta

    with app.app_context():
        # Create two students
        student1 = Student(
            control_number="TEST100",
            full_name="Student with 10+ hours",
            career="Engineering",
            email="test100@test.com",
        )
        student2 = Student(
            control_number="TEST200",
            full_name="Student with <10 hours",
            career="Engineering",
            email="test200@test.com",
        )
        db.session.add_all([student1, student2])
        db.session.flush()

        # Create activities for student1 (12 hours total)
        activity1 = Activity(
            event_id=sample_event,
            department="TEST",
            name="Activity 1",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=8),
            duration_hours=8.0,
            activity_type="Conferencia",
            location="Test",
            modality="Presencial",
        )
        activity2 = Activity(
            event_id=sample_event,
            department="TEST",
            name="Activity 2",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=4),
            duration_hours=4.0,
            activity_type="Taller",
            location="Test",
            modality="Presencial",
        )
        db.session.add_all([activity1, activity2])
        db.session.flush()

        # Register student1 with both activities (12 hours)
        reg1 = Registration(
            student_id=student1.id, activity_id=activity1.id, status="Asistió"
        )
        reg2 = Registration(
            student_id=student1.id, activity_id=activity2.id, status="Asistió"
        )

        # Create activity for student2 (5 hours only)
        activity3 = Activity(
            event_id=sample_event,
            department="TEST",
            name="Activity 3",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=5),
            duration_hours=5.0,
            activity_type="Conferencia",
            location="Test",
            modality="Presencial",
        )
        db.session.add(activity3)
        db.session.flush()

        reg3 = Registration(
            student_id=student2.id, activity_id=activity3.id, status="Asistió"
        )

        db.session.add_all([reg1, reg2, reg3])
        db.session.commit()

    # Query the endpoint
    resp = client.get(
        f"/api/students/complementary-credits?event_id={sample_event}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()

    # Should only return student1 with 12 hours
    assert "students" in data
    assert len(data["students"]) == 1
    assert data["students"][0]["control_number"] == "TEST100"
    assert data["students"][0]["total_hours"] >= 10.0


def test_export_complementary_credits_generates_excel(
    app,
    client,
    auth_headers,
    sample_event,
    sample_student,
    sample_activity,
    sample_registration,
):
    """Test that Excel export endpoint returns file."""
    resp = client.get(
        f"/api/students/complementary-credits/export?event_id={sample_event}",
        headers=auth_headers,
    )

    # Should return file or empty result depending on data
    # If no students with 10+ hours, it should still work
    assert resp.status_code == 200
    assert (
        resp.content_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_complementary_credits_include_walkins(
    app, client, auth_headers, sample_event
):
    """Los walk-ins (solo Attendance, sin Registration) sí acreditan.

    Regresión de la unificación: antes la lista solo consultaba Registrations
    con status 'Asistió' e ignoraba por completo a los estudiantes con
    check-in directo.
    """
    from app.models.attendance import Attendance
    from app.models.student import Student
    from app.models.activity import Activity
    from app import db

    with app.app_context():
        walkin = Student(
            control_number="WALK001",
            full_name="Walk In Student",
            career="Ingeniería en Sistemas",
            email="walkin@test.com",
        )
        db.session.add(walkin)
        db.session.flush()

        # 2 actividades de 5h con check-in directo (sin preregistro)
        for i in range(2):
            activity = Activity(
                event_id=sample_event,
                department="TEST",
                name=f"Walk-in Activity {i}",
                description="Test",
                start_datetime=datetime.now(),
                end_datetime=datetime.now() + timedelta(hours=5),
                duration_hours=5.0,
                activity_type="Conferencia",
                location="Test",
                modality="Presencial",
            )
            db.session.add(activity)
            db.session.flush()
            db.session.add(
                Attendance(
                    student_id=walkin.id,
                    activity_id=activity.id,
                    check_in_time=datetime.now(),
                    status="Asistió",
                    attendance_percentage=100.0,
                )
            )
        db.session.commit()

    resp = client.get(
        f"/api/students/complementary-credits?event_id={sample_event}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()

    controls = [s["control_number"] for s in data["students"]]
    assert "WALK001" in controls
    walkin_row = next(s for s in data["students"] if s["control_number"] == "WALK001")
    assert walkin_row["total_hours"] == 10.0
    assert walkin_row["has_complementary_credit"] is True


def test_complementary_credits_dedup_reg_and_attendance(
    app, client, auth_headers, sample_event
):
    """Una actividad con Registration y Attendance cuenta una sola vez."""
    from app.models.attendance import Attendance
    from app.models.student import Student
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        student = Student(
            control_number="DEDUP01",
            full_name="Dedup Student",
            career="Ingeniería en Sistemas",
            email="dedup@test.com",
        )
        db.session.add(student)
        db.session.flush()

        # Actividad de 6h con Registration 'Asistió' + Attendance 'Asistió'
        shared = Activity(
            event_id=sample_event,
            department="TEST",
            name="Shared Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=6),
            duration_hours=6.0,
            activity_type="Conferencia",
            location="Test",
            modality="Presencial",
        )
        # Actividad walk-in de 5h
        walkin = Activity(
            event_id=sample_event,
            department="TEST",
            name="Extra Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=5),
            duration_hours=5.0,
            activity_type="Taller",
            location="Test",
            modality="Presencial",
        )
        db.session.add_all([shared, walkin])
        db.session.flush()

        db.session.add(
            Registration(
                student_id=student.id, activity_id=shared.id, status="Asistió"
            )
        )
        db.session.add_all(
            [
                Attendance(
                    student_id=student.id,
                    activity_id=shared.id,
                    check_in_time=datetime.now(),
                    status="Asistió",
                    attendance_percentage=100.0,
                ),
                Attendance(
                    student_id=student.id,
                    activity_id=walkin.id,
                    check_in_time=datetime.now(),
                    status="Asistió",
                    attendance_percentage=100.0,
                ),
            ]
        )
        db.session.commit()

    resp = client.get(
        f"/api/students/complementary-credits?event_id={sample_event}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()

    row = next(
        (s for s in data["students"] if s["control_number"] == "DEDUP01"), None
    )
    assert row is not None
    # 6h (contadas una vez) + 5h walk-in = 11h (no 17h)
    assert row["total_hours"] == 11.0
    assert row["activities_count"] == 2


def test_hours_by_event_includes_walkins_alongside_registrations(
    app, client, sample_student, sample_event
):
    """Regresión del fallback parcial: los walk-ins cuentan aunque existan
    Registrations para el mismo estudiante."""
    from app.models.attendance import Attendance
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        # Registration 'Asistió' de 5h
        reg_activity = Activity(
            event_id=sample_event,
            department="TEST",
            name="Reg Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=5),
            duration_hours=5.0,
            activity_type="Conferencia",
            location="Test",
            modality="Presencial",
        )
        # Walk-in (Attendance) de 5h en otra actividad
        walkin_activity = Activity(
            event_id=sample_event,
            department="TEST",
            name="Walkin Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=5),
            duration_hours=5.0,
            activity_type="Taller",
            location="Test",
            modality="Presencial",
        )
        db.session.add_all([reg_activity, walkin_activity])
        db.session.flush()

        db.session.add(
            Registration(
                student_id=sample_student,
                activity_id=reg_activity.id,
                status="Asistió",
            )
        )
        db.session.add(
            Attendance(
                student_id=sample_student,
                activity_id=walkin_activity.id,
                check_in_time=datetime.now(),
                status="Asistió",
                attendance_percentage=100.0,
            )
        )
        db.session.commit()

    resp = client.get(f"/api/students/{sample_student}/hours-by-event")
    assert resp.status_code == 200
    data = resp.get_json()

    test_event = next(
        (e for e in data["events_hours"] if e["event_id"] == sample_event), None
    )
    assert test_event is not None
    # Antes del fallback unificado solo se veían las 5h de Registration
    assert test_event["total_hours"] == 10.0
    assert test_event["activities_count"] == 2
    assert test_event["has_complementary_credit"] is True


def test_event_details_total_and_flag_include_walkins(
    app, client, sample_student, sample_event
):
    """El total y has_complementary_credit del detalle se calculan sobre el
    total final (antes el flag se evaluaba antes de integrar walk-ins)."""
    from app.models.attendance import Attendance
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        # Registration 'Confirmado' de 8h
        reg_activity = Activity(
            event_id=sample_event,
            department="TEST",
            name="Confirmed Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=8),
            duration_hours=8.0,
            activity_type="Conferencia",
            location="Test",
            modality="Presencial",
        )
        # Walk-in de 3h
        walkin_activity = Activity(
            event_id=sample_event,
            department="TEST",
            name="Walkin Detail Activity",
            description="Test",
            start_datetime=datetime.now(),
            end_datetime=datetime.now() + timedelta(hours=3),
            duration_hours=3.0,
            activity_type="Taller",
            location="Test",
            modality="Presencial",
        )
        db.session.add_all([reg_activity, walkin_activity])
        db.session.flush()

        db.session.add(
            Registration(
                student_id=sample_student,
                activity_id=reg_activity.id,
                status="Confirmado",
            )
        )
        db.session.add(
            Attendance(
                student_id=sample_student,
                activity_id=walkin_activity.id,
                check_in_time=datetime.now(),
                status="Asistió",
                attendance_percentage=100.0,
            )
        )
        db.session.commit()

    resp = client.get(
        f"/api/students/{sample_student}/event/{sample_event}/details"
    )
    assert resp.status_code == 200
    data = resp.get_json()

    # 8h (Confirmado) + 3h (walk-in) = 11h → alcanza el crédito
    assert data["total_confirmed_hours"] == 11.0
    assert data["has_complementary_credit"] is True
    assert len(data["activities"]) == 2


# ---------------------------------------------------------------------------
# Fase 2: acumulación multi-evento (event_ids)
# ---------------------------------------------------------------------------


def _create_credits_event(app, name):
    """Crea un evento devolviendo su id."""
    from app.models.event import Event
    from app import db

    with app.app_context():
        event = Event(
            name=name,
            description="Test multi-evento",
            start_date=datetime.now(),
            end_date=datetime.now() + timedelta(days=7),
            is_active=True,
        )
        db.session.add(event)
        db.session.commit()
        return event.id


def _create_student_with_hours(app, event_hours, control, name):
    """Crea un estudiante con una actividad de X horas (status 'Asistió') por
    cada par (event_id, hours). Devuelve el student_id."""
    from app.models.student import Student
    from app.models.activity import Activity
    from app.models.registration import Registration
    from app import db

    with app.app_context():
        student = Student(
            control_number=control,
            full_name=name,
            career="Ingeniería en Sistemas",
            email=f"{control.lower()}@test.com",
        )
        db.session.add(student)
        db.session.flush()

        for idx, (ev_id, hours) in enumerate(event_hours):
            activity = Activity(
                event_id=ev_id,
                department="TEST",
                name=f"{name} Act {idx}",
                description="Test",
                start_datetime=datetime.now(),
                end_datetime=datetime.now() + timedelta(hours=hours),
                duration_hours=hours,
                activity_type="Conferencia",
                location="Test",
                modality="Presencial",
            )
            db.session.add(activity)
            db.session.flush()
            db.session.add(
                Registration(
                    student_id=student.id, activity_id=activity.id, status="Asistió"
                )
            )

        db.session.commit()
        return student.id


def test_complementary_credits_event_ids_accumulates_across_events(
    app, client, auth_headers
):
    """event_ids acumula horas entre eventos hacia el umbral de 10."""
    ev1 = _create_credits_event(app, "Aniversario 45")
    ev2 = _create_credits_event(app, "Aniversario 46")
    _create_student_with_hours(app, [(ev1, 6.0), (ev2, 4.0)], "MULT01", "Alumno Combo")
    _create_student_with_hours(
        app, [(ev1, 8.0)], "MULT02", "Alumno Insuficiente"
    )

    resp = client.get(
        f"/api/students/complementary-credits?event_ids={ev1},{ev2}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()

    # Solo el combo 6h + 4h = 10h alcanza el crédito
    assert data["total_students"] == 1
    student = data["students"][0]
    assert student["control_number"] == "MULT01"
    assert student["total_hours"] == 10.0
    assert student["hours_by_event"] == {str(ev1): 6.0, str(ev2): 4.0}
    assert student["activities_by_event"] == {str(ev1): 1, str(ev2): 1}
    assert student["has_complementary_credit"] is True

    # events incluye ambos (orden cronológico) y event es null (sin event_id)
    assert [e["id"] for e in data["events"]] == [ev1, ev2]
    assert data["event"] is None
    # Nadie queda excluido: el cruce 6+4 ocurre en el último evento
    assert data["excluded_already_credited"] == 0


def test_already_credited_in_earlier_event_is_excluded(
    app, client, auth_headers
):
    """Exclusión derivada (earliest-crossing): si la suma cruza 10h en un
    evento ANTERIOR al último seleccionado, el estudiante ya acreditado se
    omite de la lista combinada."""
    ev1 = _create_credits_event(app, "Crono Temprano")
    ev2 = _create_credits_event(app, "Crono Tardio")
    _create_student_with_hours(
        app, [(ev1, 12.0), (ev2, 3.0)], "CRON01", "Acreditado Temprano"
    )

    # Selección combinada: cruzó en ev1 (≠ último) → excluido
    resp = client.get(
        f"/api/students/complementary-credits?event_ids={ev1},{ev2}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["total_students"] == 0
    assert data["excluded_already_credited"] == 1

    # Selección solo ev1: el cruce ocurre en el único/último evento → listado
    resp = client.get(
        f"/api/students/complementary-credits?event_id={ev1}", headers=auth_headers
    )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["total_students"] == 1
    assert data["students"][0]["control_number"] == "CRON01"
    assert data["excluded_already_credited"] == 0


def test_crossing_at_intermediate_event_excluded_last_included(
    app, client, auth_headers
):
    """Con 3 eventos: el que cruza en el intermedio se excluye; el que cruza
    en el último (ejemplo del plan 4+4+2) entra."""
    ev1 = _create_credits_event(app, "Crono 1")
    ev2 = _create_credits_event(app, "Crono 2")
    ev3 = _create_credits_event(app, "Crono 3")
    # Cruza en ev2 (intermedio): 6 -> 10 -> 15
    _create_student_with_hours(
        app, [(ev1, 6.0), (ev2, 4.0), (ev3, 5.0)], "CRON02", "Cruce Intermedio"
    )
    # Cruza en ev3 (último): 4 -> 8 -> 10
    _create_student_with_hours(
        app, [(ev1, 4.0), (ev2, 4.0), (ev3, 2.0)], "CRON03", "Cruce Final"
    )

    resp = client.get(
        f"/api/students/complementary-credits?event_ids={ev1},{ev2},{ev3}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()

    assert data["total_students"] == 1
    assert data["students"][0]["control_number"] == "CRON03"
    assert data["students"][0]["total_hours"] == 10.0
    assert data["excluded_already_credited"] == 1


def test_export_event_ids_multi_event_columns(app, client, auth_headers):
    """El Excel multi-evento incluye una columna de horas por evento +
    'Horas Totales' + 'Actividades', con la misma regla que la lista."""
    from io import BytesIO

    from openpyxl import load_workbook

    ev1 = _create_credits_event(app, "Export Ev A")
    ev2 = _create_credits_event(app, "Export Ev B")
    _create_student_with_hours(
        app, [(ev1, 6.0), (ev2, 4.0)], "EXPM01", "Alumno Export Multi"
    )

    resp = client.get(
        f"/api/students/complementary-credits/export?event_ids={ev1},{ev2}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert (
        resp.content_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    headers = [cell.value for cell in ws[4]]
    assert headers == [
        "No.",
        "Número de Control",
        "Nombre Completo",
        "Carrera",
        "Email",
        "Export Ev A",
        "Export Ev B",
        "Horas Totales",
        "Actividades",
    ]
    row = [cell.value for cell in ws[5]]
    assert row[1] == "EXPM01"
    assert row[5] == 6.0  # horas en ev1
    assert row[6] == 4.0  # horas en ev2
    assert row[7] == 10.0  # total combinado
    assert row[8] == 2  # actividades
    # Solo un estudiante: la siguiente fila de datos está vacía
    assert ws.cell(row=6, column=2).value is None


def test_export_applies_derived_exclusion(app, client, auth_headers):
    """El Excel respeta la exclusión derivada (cruce en evento anterior)."""
    from io import BytesIO

    from openpyxl import load_workbook

    ev1 = _create_credits_event(app, "Export Crono 1")
    ev2 = _create_credits_event(app, "Export Crono 2")
    _create_student_with_hours(
        app, [(ev1, 12.0), (ev2, 3.0)], "EXPR01", "Alumno Ya Acreditado"
    )

    resp = client.get(
        f"/api/students/complementary-credits/export?event_ids={ev1},{ev2}",
        headers=auth_headers,
    )
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    # Encabezado presente, pero ninguna fila de datos (excluido por la regla)
    assert ws.cell(row=4, column=1).value == "No."
    assert ws.cell(row=5, column=2).value is None


def test_export_missing_event_ids_validation(client, auth_headers):
    """El export valida los mismos parámetros que la lista."""
    resp = client.get(
        "/api/students/complementary-credits/export", headers=auth_headers
    )
    assert resp.status_code == 400
    assert "event_id" in resp.get_json()["message"].lower()

    resp = client.get(
        "/api/students/complementary-credits/export?event_ids=abc",
        headers=auth_headers,
    )
    assert resp.status_code == 400

    resp = client.get(
        "/api/students/complementary-credits/export?event_ids=99999",
        headers=auth_headers,
    )
    assert resp.status_code == 404


def test_complementary_credits_event_ids_accepts_repeated_param(
    app, client, auth_headers
):
    """event_ids también acepta el parámetro repetido (?a=1&a=3)."""
    ev1 = _create_credits_event(app, "Evento Repetido A")
    ev2 = _create_credits_event(app, "Evento Repetido B")
    _create_student_with_hours(
        app, [(ev1, 5.0), (ev2, 5.0)], "REP01", "Alumno Repetido"
    )

    resp = client.get(
        f"/api/students/complementary-credits?event_ids={ev1}&event_ids={ev2}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["total_students"] == 1
    assert data["students"][0]["total_hours"] == 10.0


def test_complementary_credits_merges_event_id_and_event_ids(
    app, client, auth_headers
):
    """Enviar event_id y event_ids juntos produce la unión de ambos."""
    ev1 = _create_credits_event(app, "Evento Union A")
    ev2 = _create_credits_event(app, "Evento Union B")
    _create_student_with_hours(
        app, [(ev1, 7.0), (ev2, 3.0)], "UNI01", "Alumno Union"
    )

    resp = client.get(
        f"/api/students/complementary-credits?event_id={ev1}&event_ids={ev2}",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.get_json()

    # Contrato original: 'event' sigue siendo el payload del event_id
    assert data["event"] is not None
    assert data["event"]["id"] == ev1
    assert "name" in data["event"]
    assert "start_date" in data["event"]
    assert "end_date" in data["event"]
    # Aditivo: 'events' trae la unión
    assert [e["id"] for e in data["events"]] == [ev1, ev2]
    assert data["total_students"] == 1
    assert data["students"][0]["total_hours"] == 10.0


def test_complementary_credits_single_event_keeps_original_contract(
    app, client, auth_headers
):
    """El modo original (solo event_id) conserva su respuesta e incluye
    events con un único elemento."""
    ev = _create_credits_event(app, "Evento Original")
    _create_student_with_hours(app, [(ev, 12.0)], "OLD01", "Alumno Viejo")

    resp = client.get(
        f"/api/students/complementary-credits?event_id={ev}", headers=auth_headers
    )
    assert resp.status_code == 200
    data = resp.get_json()

    assert data["event"]["id"] == ev
    assert data["event"]["name"] == "Evento Original"
    assert data["event"]["start_date"] is not None
    assert data["event"]["end_date"] is not None
    assert len(data["events"]) == 1
    assert data["events"][0]["id"] == ev
    assert data["total_students"] == 1
    assert data["students"][0]["control_number"] == "OLD01"
    # hours_by_event también está disponible en modo single (aditivo)
    assert data["students"][0]["hours_by_event"] == {str(ev): 12.0}


def test_complementary_credits_event_ids_validation(client, auth_headers):
    """Validaciones del nuevo parámetro."""
    # Sin parámetros: mensaje original intacto
    resp = client.get("/api/students/complementary-credits", headers=auth_headers)
    assert resp.status_code == 400
    assert "event_id" in resp.get_json()["message"].lower()

    # IDs no numéricos
    resp = client.get(
        "/api/students/complementary-credits?event_ids=abc", headers=auth_headers
    )
    assert resp.status_code == 400
    assert "numéricos" in resp.get_json()["message"]

    # Evento inexistente: mismo 404 de siempre
    resp = client.get(
        "/api/students/complementary-credits?event_ids=99999", headers=auth_headers
    )
    assert resp.status_code == 404
    assert resp.get_json()["message"] == "Evento no encontrado"

    # Mezcla válida + inválida también rechaza
    resp = client.get(
        "/api/students/complementary-credits?event_ids=1,xyz", headers=auth_headers
    )
    assert resp.status_code == 400
